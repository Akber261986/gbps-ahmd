from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def generate_students_excel(students, school_name: str, semis_code: str):
    """
    Generate a professionally formatted Excel workbook
    containing active students.
    """

    wb = Workbook()
    ws = wb.active
    ws.title = "Student Register"

    # ---------------------------------------------------------
    # COLORS / STYLES
    # ---------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    title_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
        size=11
    )

    title_font = Font(
        bold=True,
        size=16
    )

    subtitle_font = Font(
        bold=True,
        size=12
    )

    normal_font = Font(
        size=11
    )

    thin_side = Side(
        style="thin",
        color="B7B7B7"
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    center_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    left_alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    # ---------------------------------------------------------
    # SCHOOL HEADER
    # ---------------------------------------------------------

    ws.merge_cells("A1:G1")
    ws["A1"] = school_name
    ws["A1"].font = title_font
    ws["A1"].alignment = center_alignment
    ws["A1"].fill = title_fill

    ws.merge_cells("A2:G2")
    ws["A2"] = "ACTIVE STUDENTS REGISTER"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = center_alignment
    ws["A2"].fill = title_fill

    ws.merge_cells("A3:C3")
    ws["A3"] = f"SEMIS CODE: {semis_code}"
    ws["A3"].font = Font(bold=True, size=10)

    ws.merge_cells("E3:G3")
    ws["E3"] = f"Total Students: {len(students)}"
    ws["E3"].font = Font(bold=True, size=10)
    ws["E3"].alignment = Alignment(horizontal="right")

    # ---------------------------------------------------------
    # TABLE HEADERS
    # ---------------------------------------------------------

    header_row = 5

    headers = [
        "S.No",
        "G.R No",
        "Student Name",
        "Father's Name",
        "Class",
        "Medium",
        "Religion",
    ]

    for col_num, header in enumerate(headers, start=1):
        cell = ws.cell(
            row=header_row,
            column=col_num,
            value=header
        )

        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center_alignment
        cell.border = border

    # ---------------------------------------------------------
    # STUDENT DATA
    # ---------------------------------------------------------

    for index, student in enumerate(students, start=1):

        row = header_row + index

        values = [
            index,
            student.gr_no,
            student.name,
            student.father_name,
            student.class_name,
            student.medium,
            student.religion,
        ]

        for col_num, value in enumerate(values, start=1):

            cell = ws.cell(
                row=row,
                column=col_num,
                value=value
            )

            cell.font = normal_font
            cell.border = border

            # Numbers centered
            if col_num in [1, 2, 5]:
                cell.alignment = center_alignment
            else:
                cell.alignment = left_alignment

    # ---------------------------------------------------------
    # COLUMN WIDTHS
    # ---------------------------------------------------------

    column_widths = {
        "A": 8,
        "B": 12,
        "C": 28,
        "D": 28,
        "E": 12,
        "F": 15,
        "G": 15,
    }

    for column, width in column_widths.items():
        ws.column_dimensions[column].width = width

    # ---------------------------------------------------------
    # ROW HEIGHTS
    # ---------------------------------------------------------

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[5].height = 24

    # ---------------------------------------------------------
    # FREEZE HEADER
    # ---------------------------------------------------------

    ws.freeze_panes = "A6"

    # ---------------------------------------------------------
    # AUTO FILTER
    # ---------------------------------------------------------

    if students:
        last_row = header_row + len(students)

        ws.auto_filter.ref = f"A5:G{last_row}"

    # ---------------------------------------------------------
    # EXCEL TABLE
    # ---------------------------------------------------------

    if students:
        last_row = header_row + len(students)

        table_ref = f"A5:G{last_row}"

        table = Table(
            displayName="ActiveStudents",
            ref=table_ref
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = style

        ws.add_table(table)

    # ---------------------------------------------------------
    # PRINT SETTINGS
    # ---------------------------------------------------------

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.print_title_rows = "1:5"

    ws.sheet_view.showGridLines = False

    # ---------------------------------------------------------
    # PAGE MARGINS
    # ---------------------------------------------------------

    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    # ---------------------------------------------------------
    # SAVE TO MEMORY
    # ---------------------------------------------------------

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return output
